import pandas as pd
import warnings
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Alignment

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# 四舍六入五成双函数
def round_half_even(value, decimals=2):
    if pd.isna(value):
        return value
    multiplier = 10 ** decimals
    value *= multiplier
    floor_val = int(value)
    fractional = value - floor_val
    
    if fractional < 0.5:
        result = floor_val
    elif fractional > 0.5:
        result = floor_val + 1
    else:
        if floor_val % 2 == 0:
            result = floor_val
        else:
            result = floor_val + 1
    
    return result / multiplier

# ----------------------------- 1. 文件上传 -----------------------------
root = tk.Tk()
root.withdraw()  # 隐藏主窗口

file_path = filedialog.askopenfilename(
    title="请选择权责收入表文件",
    filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
)

if not file_path:
    messagebox.showwarning("提示", "未选择文件，程序退出。")
    exit()

df_raw = pd.read_excel(file_path, header=[0, 1])

# 扁平化列名（与之前相同）
new_columns = []
for col in df_raw.columns:
    level0 = str(col[0]).strip()
    level1 = str(col[1]).strip()
    if level1 in ("nan", "") or level1.startswith("Unnamed"):
        new_columns.append(level0)
    else:
        new_columns.append(f"{level0}_{level1}")
df_raw.columns = new_columns

# 去掉“合计”行
target_col = "序号"
if target_col in df_raw.columns:
    df_raw = df_raw[~df_raw[target_col].astype(str).str.contains("合计", na=False)]

# 长表转换（月份列格式为“2026年01月”）
month_cols = [c for c in df_raw.columns if "年" in c and ("_含税" in c or "_不含税" in c)]
months_orig = sorted(set([c.split("_")[0] for c in month_cols]))

rows = []
for _, row in df_raw.iterrows():
    for month in months_orig:
        tax_col = f"{month}_含税"
        net_col = f"{month}_不含税"
        if tax_col not in df_raw.columns or net_col not in df_raw.columns:
            continue
        tax_val = row[tax_col]
        net_val = row[net_col]
        if pd.isna(tax_val) and pd.isna(net_val):
            continue
        rows.append({
            "品牌": row["品牌"],
            "铺位/点位号": row["铺位/点位号"],
            "费用名称": row["费用名称"],
            "月份": month,
            "含税收入": tax_val,
            "不含税收入": net_val
        })

df_long = pd.DataFrame(rows)

# 转换月份为显示格式“2026-01”
def convert_month_format(orig):
    return f"{orig[:4]}-{orig[5:7]}"

df_long["月份_显示"] = df_long["月份"].apply(convert_month_format)
unique_months_display = sorted(df_long["月份_显示"].unique())

if len(unique_months_display) == 0:
    messagebox.showerror("错误", "没有找到任何月份数据，请检查源文件。")
    exit()

# ----------------------------- 2. 输入框获取需要导出的月份 -----------------------------
# 显示可用月份提示
available_months_str = "、".join(unique_months_display)
prompt = f"当前数据中包含的月份：\n{available_months_str}\n\n请输入需要导出的月份（多个用英文逗号分隔）：\n例如：2026-01,2026-03"

user_input = simpledialog.askstring("选择月份", prompt, initialvalue="")

if not user_input:
    messagebox.showwarning("提示", "未输入月份，程序退出。")
    exit()

# 解析输入的月份
selected_months_display = [m.strip() for m in user_input.split(",") if m.strip() != ""]
# 去除重复并保持输入顺序（但后续排序也可以）
selected_months_display = list(dict.fromkeys(selected_months_display))  # 去重保留顺序

# 验证输入的月份是否在可用月份中
invalid_months = [m for m in selected_months_display if m not in unique_months_display]
if invalid_months:
    messagebox.showerror("错误", f"以下月份不在数据中：{', '.join(invalid_months)}\n可用月份：{available_months_str}")
    exit()

# 将显示格式映射回原始格式（2026-01 → 2026年01月）
def display_to_orig(display):
    year, month = display.split("-")
    return f"{year}年{month}月"

selected_months_orig = [display_to_orig(m) for m in selected_months_display]

# 过滤数据
df_filtered = df_long[df_long["月份"].isin(selected_months_orig)].copy()
if df_filtered.empty:
    messagebox.showerror("错误", "所选月份在数据中不存在，请检查。")
    exit()

# 删除辅助列，并将月份列转换为显示格式
df_filtered.drop(columns=["月份_显示"], inplace=True)
df_filtered["月份"] = df_filtered["月份"].apply(convert_month_format)

# 仅保留指定的六个字段
df_filtered = df_filtered[["品牌", "铺位/点位号", "费用名称", "月份", "含税收入", "不含税收入"]]

# ----------------------------- 3. 选择保存文件夹并按费用名称拆分 -----------------------------
save_dir = filedialog.askdirectory(title="请选择保存文件夹")
if not save_dir:
    messagebox.showwarning("提示", "未选择文件夹，程序退出。")
    exit()

output_path = f"{save_dir}/权责收入表_按月费拆分.xlsx"

grouped = df_filtered.groupby("费用名称")
sheet_names = list(grouped.groups.keys())

if not sheet_names:
    messagebox.showerror("错误", "没有有效的费用名称，无法拆分。")
    exit()

# ========== 步骤1: 将数据拆分为多个工作表 ==========
with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    for fee_name, data in grouped:
        safe_name = str(fee_name).strip()
        invalid_chars = r'[]:*?/\\'
        for ch in invalid_chars:
            safe_name = safe_name.replace(ch, '_')
        if len(safe_name) > 31:
            safe_name = safe_name[:31]
        if safe_name == "":
            safe_name = "未命名费用"
        data.to_excel(writer, sheet_name=safe_name, index=False)

# 创建边框样式
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 加载工作簿进行后续格式设置
wb = load_workbook(output_path)

# 遍历每个工作表执行后续操作
for fee_name in sheet_names:
    safe_name = str(fee_name).strip()
    invalid_chars = r'[]:*?/\\'
    for ch in invalid_chars:
        safe_name = safe_name.replace(ch, '_')
    if len(safe_name) > 31:
        safe_name = safe_name[:31]
    if safe_name == "":
        safe_name = "未命名费用"
    
    ws = wb[safe_name]
    max_row = ws.max_row
    max_col = ws.max_column
    
    # ========== 步骤2: 对数值数据应用四舍六入五成双规则 ==========
    for row in range(2, max_row + 1):
        for col in range(5, max_col + 1):  # E列和F列（含税收入和不含税收入）
            cell = ws.cell(row=row, column=col)
            if cell.value is not None:
                try:
                    cell.value = round_half_even(float(cell.value), 2)
                except:
                    pass
    
    # ========== 步骤3: 添加合计行 ==========
    total_tax = 0
    total_net = 0
    for row in range(2, max_row + 1):
        tax_cell = ws.cell(row=row, column=5)
        net_cell = ws.cell(row=row, column=6)
        if tax_cell.value is not None:
            try:
                total_tax += float(tax_cell.value)
            except:
                pass
        if net_cell.value is not None:
            try:
                total_net += float(net_cell.value)
            except:
                pass
    
    total_row = max_row + 1
    ws.cell(row=total_row, column=1, value="合计")
    ws.cell(row=total_row, column=2, value="")
    ws.cell(row=total_row, column=3, value="")
    ws.cell(row=total_row, column=4, value="")
    ws.cell(row=total_row, column=5, value=total_tax)
    ws.cell(row=total_row, column=6, value=total_net)
    
    # ========== 步骤4: 为所有包含数据的单元格添加边框 ==========
    for row in ws.iter_rows(min_row=1, max_row=total_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin_border
    
    # ========== 步骤5: 合并合计行的A列至D列 ==========
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal='center', vertical='center')

# 保存最终结果
wb.save(output_path)

messagebox.showinfo("完成", f"已成功按费用名称拆分数据，共 {len(sheet_names)} 个费用类型。\n保存位置：{output_path}")
print(f"处理完成！文件保存至：{output_path}")